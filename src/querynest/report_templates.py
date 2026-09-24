"""Template reports with REAL Excel PivotTables.

Python libraries can't create a PivotTable, but openpyxl can keep one that already exists.
So we split the work:

  1. Authoring (once, on a Windows machine with Excel):  uv run make-templates
     Excel itself builds report_templates/<name>.xlsx: a "Data" sheet holding an Excel Table
     named SourceData, and a "Pivot" sheet with a real PivotTable on that table, set to refresh
     when the file is opened. Only placeholder rows go in: a template never contains real data.
  2. Filling (every request, on the server, no Excel):  fill_template()
     openpyxl opens the template, writes this user's rows into SourceData, resizes the table,
     drops the pivot's cached placeholder records and marks the cache "refresh on load".
     When the user opens the file, Excel rebuilds the pivot from the real rows.

Templates are defined in report_templates/templates.json: title, description, a fixed SQL query
(written by an admin, still run with the user's permissions and guardrails) and the pivot layout.
"""

import io
import json
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from querynest.config import BACKEND_DIR

TEMPLATE_DIR = BACKEND_DIR / "report_templates"
DATA_SHEET, PIVOT_SHEET, TABLE_NAME = "Data", "Pivot", "SourceData"
EXCEL_AGG = {"sum": -4157, "count": -4112, "avg": -4106, "min": -4139, "max": -4136}  # xlSum, xlCount...


@dataclass
class PivotLayout:
    rows: list[str]
    columns: list[str] = field(default_factory=list)
    filters: list[str] = field(default_factory=list)
    values: list[dict[str, str]] = field(default_factory=list)  # [{"field": "revenue", "agg": "sum"}]


@dataclass
class Template:
    name: str
    title: str
    description: str
    sql: str
    pivot: PivotLayout

    @property
    def path(self) -> Path:
        return TEMPLATE_DIR / f"{self.name}.xlsx"

    @property
    def ready(self) -> bool:
        return self.path.exists()


def load_templates() -> dict[str, Template]:
    raw = json.loads((TEMPLATE_DIR / "templates.json").read_text(encoding="utf-8"))
    return {name: Template(name, t["title"], t["description"], t["sql"], PivotLayout(**t["pivot"]))
            for name, t in raw.items()}


# ---------------------------------------------------------------------------
# 1. Authoring with Excel (Windows + Excel + pywin32; dev machines only)
# ---------------------------------------------------------------------------
def placeholder_rows(columns: list[str], kinds: dict[str, str]) -> list[list[Any]]:
    """Two fake rows with the right types, so Excel knows which fields are numbers."""
    def value(col: str, i: int) -> Any:
        kind = kinds.get(col, "text")
        return i if kind == "number" else datetime(2000, 1, i) if kind == "date" else f"(sample {i})"
    return [[value(c, i) for c in columns] for i in (1, 2)]


RPC_E_CALL_REJECTED = -2147418111  # "Call was rejected by callee": Excel is busy right now
_excel_pid: int | None = None  # the Excel process we're automating (for dialog checks + cleanup)


class ExcelBlocked(RuntimeError):
    def __init__(self, dialog: str):
        super().__init__(f"Excel is waiting on its dialog '{dialog}', which automation can't answer. "
                         "Open Excel normally once, answer that dialog, close Excel, then run "
                         "`uv run make-templates` again.")


def _retry(action, seconds: float = 60):
    """Run one Excel COM operation, retrying while Excel reports it is busy."""
    import time

    import pythoncom
    import pywintypes

    deadline = time.monotonic() + seconds
    while True:
        try:
            return action()
        except pywintypes.com_error as e:
            if e.hresult != RPC_E_CALL_REJECTED or time.monotonic() > deadline:
                raise
            if _excel_pid and time.monotonic() > deadline - seconds + 3:
                if (dialog := _blocking_dialog(_excel_pid)) is not None:
                    raise ExcelBlocked(dialog) from None
            # Excel may be waiting on OUR thread's message queue; process it, then retry
            pythoncom.PumpWaitingMessages()
            time.sleep(0.2)


def _unwrap(value):
    return object.__getattribute__(value, "_obj") if isinstance(value, _Patient) else value


class _Patient:
    """Proxy around an Excel COM object: every attribute read, write and call is retried while
    Excel is busy, and objects it returns are wrapped too (excel.Workbooks.Add() etc.)."""

    def __init__(self, obj):
        object.__setattr__(self, "_obj", obj)

    @staticmethod
    def _wrap(value):
        return _Patient(value) if hasattr(value, "_oleobj_") or callable(value) else value

    def __getattr__(self, name):
        return self._wrap(_retry(lambda: getattr(object.__getattribute__(self, "_obj"), name)))

    def __setattr__(self, name, value):
        _retry(lambda: setattr(object.__getattribute__(self, "_obj"), name, _unwrap(value)))

    def __call__(self, *args, **kwargs):
        args = [_unwrap(a) for a in args]
        kwargs = {k: _unwrap(v) for k, v in kwargs.items()}
        return self._wrap(_retry(lambda: object.__getattribute__(self, "_obj")(*args, **kwargs)))


def _blocking_dialog(excel_pid: int) -> str | None:
    """Title of a visible Office dialog (e.g. the first-run 'Your privacy option') in this Excel
    process. Invisible Excel can't answer it, so every COM call would be rejected forever."""
    import win32gui
    import win32process

    found: list[str] = []

    def check(hwnd, _):
        if (win32process.GetWindowThreadProcessId(hwnd)[1] == excel_pid and win32gui.IsWindowVisible(hwnd)
                and win32gui.GetClassName(hwnd) in ("NUIDialog", "#32770")):
            found.append(win32gui.GetWindowText(hwnd))
        return True

    win32gui.EnumWindows(check, None)
    return found[0] if found else None


def author_template(template: Template, columns: list[str], kinds: dict[str, str]) -> None:
    global _excel_pid
    import subprocess

    import pythoncom
    import win32com.client as win32
    import win32process

    pythoncom.CoInitialize()
    raw = win32.DispatchEx("Excel.Application")  # a private, invisible Excel instance
    _excel_pid = win32process.GetWindowThreadProcessId(raw.Hwnd)[1]
    excel = _Patient(raw)
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Add()
        data = wb.Worksheets(1)
        data.Name = DATA_SHEET
        rows = [columns, *placeholder_rows(columns, kinds)]
        cells = data.Range(data.Cells(1, 1), data.Cells(len(rows), len(columns)))
        cells.Value = rows
        table = data.ListObjects.Add(1, cells, None, 1)  # xlSrcRange, headers: xlYes
        table.Name = TABLE_NAME
        table.TableStyle = "TableStyleMedium2"

        pivot_ws = wb.Worksheets.Add()
        pivot_ws.Name = PIVOT_SHEET
        pivot_ws.Range("A1").Value = template.title
        pivot_ws.Range("A1").Font.Bold = True
        pivot_ws.Range("A1").Font.Size = 16
        cache = wb.PivotCaches().Create(SourceType=1, SourceData=TABLE_NAME)  # xlDatabase, by table name
        pivot = cache.CreatePivotTable(TableDestination=pivot_ws.Range("A4"), TableName="ReportPivot")
        layout = template.pivot
        for name in layout.filters:
            pivot.PivotFields(name).Orientation = 3  # xlPageField (filter at the top)
        for name in layout.rows:
            pivot.PivotFields(name).Orientation = 1  # xlRowField
        for name in layout.columns:
            pivot.PivotFields(name).Orientation = 2  # xlColumnField
        for v in layout.values:
            data_field = pivot.AddDataField(pivot.PivotFields(v["field"]), f"{v['agg'].title()} of {v['field']}",
                                            EXCEL_AGG[v["agg"]])
            data_field.NumberFormat = "#,##0.000"
        pivot.PivotCache().RefreshOnFileOpen = True
        pivot_ws.Activate()
        TEMPLATE_DIR.mkdir(exist_ok=True)
        wb.SaveAs(str(template.path.resolve()), 51)  # 51 = .xlsx
        wb.Close(False)
    finally:
        try:
            raw.Quit()
        except Exception:  # never let a failed Quit hide the real error above
            pass
        # Make sure OUR hidden Excel is gone (a stuck instance would block the next run)
        subprocess.run(["taskkill", "/PID", str(_excel_pid), "/F"], capture_output=True)
        _excel_pid = None
        pythoncom.CoUninitialize()


# ---------------------------------------------------------------------------
# 2. Filling on the server (openpyxl, no Excel)
# ---------------------------------------------------------------------------
def fill_template(template: Template, columns: list[str], rows: list[dict[str, Any]],
                  kinds: dict[str, str], author: str) -> bytes:
    if not template.ready:
        raise FileNotFoundError(f"Template '{template.name}' has not been authored yet (run make-templates)")
    wb = openpyxl.load_workbook(template.path)  # keeps the PivotTable parts intact
    ws = wb[DATA_SHEET]
    table = ws.tables[TABLE_NAME]
    header = [c.value for c in ws[1]][: len(columns)]
    if header != columns:
        raise ValueError(f"Template columns {header} don't match the query's {columns}; re-run make-templates")

    ws.delete_rows(2, ws.max_row)  # remove placeholder rows
    for r in rows:
        ws.append([_excel_value(r.get(c), kinds.get(c, "text")) for c in columns])
    for i, kind in enumerate((kinds.get(c, "text") for c in columns), start=1):
        fmt = "#,##0.000" if kind == "number" else "yyyy-mm-dd" if kind == "date" else None
        if fmt:
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = fmt
    last_col = openpyxl.utils.get_column_letter(len(columns))
    table.ref = f"A1:{last_col}{max(len(rows), 1) + 1}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref

    pivot_ws = wb[PIVOT_SHEET]
    pivot_ws["A2"] = f"{len(rows):,} rows · generated {datetime.now():%Y-%m-%d %H:%M} for {author}"
    for sheet in wb.worksheets:
        for pivot in sheet._pivots:
            cache = pivot.cache
            cache.refreshOnLoad = True  # Excel rebuilds the pivot from SourceData when opened
            cache.records = None  # drop cached placeholder rows so no stale data ships
            cache.saveData = False
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _excel_value(value: Any, kind: str) -> Any:
    if value is None:
        return None
    if kind == "date" and isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return value
    return value
