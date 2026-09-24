"""The four SQLBot-inspired features: schema pre-selection, fast path, follow-ups, template reports."""

import json

import openpyxl
import pytest
from conftest import needs_db
from fakes import FakeProvider

from querynest import schema
from querynest.agent import Agent, parse_json
from querynest.config import settings
from querynest.knowledge import retrieve
from querynest.llm import LLMRouter
from querynest.masking import Masker
from querynest.permissions import UserContext, system_user

pytestmark = needs_db
MANAGER = system_user("manager")
SALES = UserContext(username="t-sales", role="sales", attributes={"salesman": "Salesman 12"})
SQL = ("SELECT mname AS category, SUM(salesamt) AS revenue FROM sales.invoices WHERE cyyyy = 2025 "
       "GROUP BY 1 ORDER BY 2 DESC")


def plan(**fields):
    return ("text", json.dumps(fields))


def run(steps, user=MANAGER, **overrides):
    fake = FakeProvider(steps)
    agent = Agent(user, LLMRouter([fake], use_cache=False), Masker())
    return agent.run("Revenue by category in 2025 as a chart"), fake


# ------------------------------------------------------------------ (c) schema pre-selection
def test_schema_lists_only_visible_tables_and_columns():
    schema.clear_cache()
    manager_tables = {t.name for t in schema.visible_tables(MANAGER)}
    sales_tables = schema.visible_tables(SALES)
    assert manager_tables == {"sales.invoices", "sales.account_entries"}
    assert [t.name for t in sales_tables] == ["sales.invoices"]
    assert "costvalue" not in {c.name for c in sales_tables[0].columns}


def test_schema_ranking_prefers_the_relevant_table():
    question = "How much VAT did we collect per year?"
    top = schema.select_tables(question, MANAGER, retrieve(question), k=1)
    assert top[0].name == "sales.account_entries"  # the VAT glossary term points to the ledger


def test_schema_prompt_has_column_comments():
    text = schema.schema_prompt(schema.visible_tables(MANAGER))
    assert "# sales.invoices" in text and "salesmna (text): Salesman name" in text


# ------------------------------------------------------------------ (d) fast path
def test_fast_path_answers_in_two_calls_without_tools(monkeypatch):
    monkeypatch.setattr(settings, "fast_path", True)
    out, fake = run([plan(mode="simple", sql=SQL, chart="bar", title="Revenue by category"),
                     ("text", "Equipment sales lead.")])
    assert out["path"] == "fast" and len(fake.requests) == 2
    assert all(r["tools"] == [] for r in fake.requests)  # no tool schemas sent at all
    assert [b["kind"] for b in out["blocks"]] == ["result", "chart"]
    assert "Equipment" in out["text"]


def test_fast_path_answer_call_gets_masked_data(monkeypatch):
    monkeypatch.setattr(settings, "fast_path", True)
    sql = ("SELECT customername, SUM(salesamt) AS revenue FROM sales.invoices WHERE cyyyy = 2025 "
           "GROUP BY 1 ORDER BY 2 DESC LIMIT 3")
    out, fake = run([plan(mode="simple", sql=sql, chart="none", title="Top 3"), ("text", "Top is [P1].")])
    answer_request = json.dumps(fake.requests[1]["messages"], default=str)
    assert "Customer 0352" not in answer_request and "[P1]" in answer_request
    assert "Customer 0352" in out["text"]


def test_fast_path_falls_back_to_agent(monkeypatch):
    monkeypatch.setattr(settings, "fast_path", True)
    # The plan asks for the agent; the agent then runs a query and answers
    out, fake = run([plan(mode="agent"), ("tool", "run_sql_query", {"sql": SQL}), ("text", "done")])
    assert out["path"] == "agent" and len(fake.requests) == 3
    assert fake.requests[1]["tools"]  # the agent call has tools


@pytest.mark.parametrize("bad_plan", [("text", "not json at all"),
                                      plan(mode="simple", sql="DELETE FROM sales.invoices")])
def test_bad_or_blocked_plan_falls_back(monkeypatch, bad_plan):
    monkeypatch.setattr(settings, "fast_path", True)
    out, _ = run([bad_plan, ("tool", "run_sql_query", {"sql": SQL}), ("text", "done")])
    assert out["path"] == "agent" and out["text"] == "done"


def test_parse_json_tolerates_fences_and_chatter():
    assert parse_json('Sure!\n```json\n{"mode": "agent"}\n```') == {"mode": "agent"}
    assert parse_json('["a", "b"]') == ["a", "b"]
    assert parse_json("nope") is None


# ------------------------------------------------------------------ greetings
@pytest.mark.parametrize("text", ["hi", "Hello!", "hey there", "Good morning QueryNest", "hii 👋"])
def test_greeting_is_answered_without_llm(text):
    from querynest.agent import GREETING_REPLY

    fake = FakeProvider([])  # any LLM call would fail: no scripted steps
    out = Agent(MANAGER, LLMRouter([fake], use_cache=False)).run(text)
    assert out["text"] == GREETING_REPLY and fake.requests == []


@pytest.mark.parametrize("text", ["hi, what was revenue in 2025?", "hello show top customers", "history of sales"])
def test_questions_starting_with_hi_are_not_greetings(text):
    from querynest.agent import is_greeting

    assert not is_greeting(text)


# ------------------------------------------------------------------ (b) follow-up suggestions
def test_followup_suggestions_block(monkeypatch):
    monkeypatch.setattr(settings, "fast_path", True)
    monkeypatch.setattr(settings, "suggest_followups", True)
    out, fake = run([plan(mode="simple", sql=SQL, chart="none", title="x"), ("text", "Answer."),
                     ("text", '["Same for 2024?", "Which brand leads?", "Monthly trend?"]')])
    assert out["blocks"][-1] == {"kind": "suggestions", "items": ["Same for 2024?", "Which brand leads?", "Monthly trend?"]}
    assert "revenue" in fake.requests[-1]["messages"][0]["content"]  # column names only
    assert "Data:" not in fake.requests[-1]["messages"][0]["content"]  # no rows are sent


# ------------------------------------------------------------------ (a) template reports
def test_template_report_file_has_real_pivot_and_only_the_users_rows():
    from querynest.report_templates import fill_template, load_templates
    from querynest.api.routes import template_report
    from querynest.appdb import User, session
    from sqlalchemy import select

    template = load_templates()["sales_pivot"]
    if not template.ready:
        pytest.skip("sales_pivot.xlsx not authored (uv run make-templates)")
    with session() as s:
        sales_user = s.scalars(select(User).where(User.username == "sales12")).first()
    import io
    wb = openpyxl.load_workbook(io.BytesIO(template_report("sales_pivot", sales_user)))
    rows = list(wb["Data"].iter_rows(min_row=2, values_only=True))
    assert rows and {r[2] for r in rows} == {"Salesman 12"}  # row filter applies to reports too
    assert wb["Data"].tables["SourceData"].ref == f"A1:I{len(rows) + 1}"
    pivot = wb["Pivot"]._pivots[0]
    assert pivot.cache.refreshOnLoad and pivot.cache.records is None  # no stale cached data ships
    assert "(sample" not in str(rows)  # placeholders are gone
    with pytest.raises(ValueError, match="don't match"):
        fill_template(template, ["wrong"], [], {}, "x")
