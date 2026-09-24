"""M6-M8: charts, pivots and Excel (no database needed)."""

import io

import openpyxl
import pytest

from querynest import charts, excel, pivots
from querynest.appdb import ResultSet


def result(rows, types, title="Test"):
    return ResultSet(id="r1", title=title, sql="SELECT 1", columns=list(types), column_types=types,
                     rows=rows, row_count=len(rows), truncated=False)


MONTHS = result([{"month": f"2025-{m:02d}-01", "revenue": m * 10.0} for m in range(1, 13)],
                {"month": "date", "revenue": "number"})
SHARES = result([{"category": c, "revenue": v} for c, v in [("A", 5.0), ("B", 3.0), ("C", 2.0)]],
                {"category": "text", "revenue": "number"})


def test_auto_chart_picks_line_for_time_and_pie_for_shares():
    assert charts.build_chart(MONTHS)["chart_type"] == "line"
    assert charts.build_chart(SHARES)["chart_type"] == "pie"


def test_chart_sums_rows_that_share_an_x_value():
    detailed = result([{"cat": c, "q": q, "v": 1.0} for c in "AB" for q in (1, 2, 3)],
                      {"cat": "text", "q": "number", "v": "number"})
    spec = charts.build_chart(detailed, "bar", "cat", ["v"])
    assert spec["data"] == [{"cat": "A", "v": 3.0}, {"cat": "B", "v": 3.0}] and spec["summed"]


def test_chart_rejects_text_measure():
    with pytest.raises(charts.ChartError, match="not numeric"):
        charts.build_chart(SHARES, "bar", "revenue", ["category"])


def test_pivot_totals_and_natural_order():
    data = result([{"b": b, "m": m, "v": 1.0} for b in "XY" for m in (1, 2, 10)], {"b": "text", "m": "number", "v": "number"})
    spec = pivots.build_pivot(data, ["b"], "v", "m", "sum")
    assert spec["header"] == ["b", "1", "2", "10", "Total"]
    assert spec["rows"][-1] == ["Total", 2.0, 2.0, 2.0, 6.0]


def test_pivot_rejects_unknown_column():
    with pytest.raises(pivots.PivotError, match="not in the result"):
        pivots.build_pivot(SHARES, ["nope"], "revenue")


def test_excel_report_has_cover_data_chart_and_formula_pivot():
    spec = pivots.build_pivot(SHARES, ["category"], "revenue", None, "sum")
    chart = charts.build_chart(SHARES)
    data = excel.build_report("My report", "tester", [{"result": SHARES, "charts": [chart], "pivots": [spec]}])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Cover" and len(wb.sheetnames) == 3
    assert wb.sheetnames[1] in wb["Cover"]["A6"].value
    pivot = wb[wb.sheetnames[2]]
    assert str(pivot["B5"].value).startswith("=SUMIFS(")
    assert len(wb[wb.sheetnames[1]]._charts) == 1


def test_filename_is_safe():
    assert excel.filename("Revenue: 2025/Q1 <final>") == "Revenue_2025Q1_final.xlsx"
